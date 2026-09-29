#!/usr/bin/env python
"""Gera um .docx a partir de um .md (mesmo fluxo usado em arquitetura_tecnica.docx).

Suporta: títulos, parágrafos, listas, tabelas, blocos de código (monoespaçado,
preservando diagramas ASCII), citações e formatação inline (**negrito**, `código`,
*itálico*, [texto](link)).

Uso:
    python md_to_docx.py planejamento_tg.md planejamento_tg.docx
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CODE_FONT = "Consolas"
CODE_SIZE = Pt(7)
CODE_SHADING = "F2F2F2"

INLINE_RE = re.compile(
    r"(\*\*.+?\*\*"      # negrito
    r"|`[^`]+`"          # código inline
    r"|\*[^*\n]+?\*"     # itálico
    r"|\[[^\]]+\]\([^)]+\))"  # link
)


def add_inline(paragraph, text: str) -> None:
    """Escreve `text` no parágrafo interpretando a formatação inline do Markdown."""
    for token in INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(9)
        elif token.startswith("*") and token.endswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        elif token.startswith("["):
            label, _, url = token[1:].partition("](")
            run = paragraph.add_run(label)
            run.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)
            run.underline = True
        else:
            paragraph.add_run(token)


def add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line if line else " ")
        run.font.name = CODE_FONT
        run.font.size = CODE_SIZE
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), CODE_SHADING)
        p._p.get_or_add_pPr().append(shd)


def add_table(doc: Document, rows: list[str]) -> None:
    def cells(row: str) -> list[str]:
        return [c.strip() for c in row.strip().strip("|").split("|")]

    header = cells(rows[0])
    body = [cells(r) for r in rows[2:]]  # rows[1] é a linha de separação (|---|)

    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        cell.paragraphs[0].text = ""
        add_inline(cell.paragraphs[0], text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)

    for row_data in body:
        cells_out = table.add_row().cells
        for cell, text in zip(cells_out, row_data):
            cell.paragraphs[0].text = ""
            add_inline(cell.paragraphs[0], text)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
    doc.add_paragraph()


def convert(md_path: str, docx_path: str) -> None:
    lines = open(md_path, encoding="utf-8").read().split("\n")
    doc = Document()

    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    i = 0
    stats = {"headings": 0, "tables": 0, "code": 0, "paragraphs": 0}

    # O Markdown quebra parágrafos em várias linhas; só linha em branco (ou um
    # elemento estrutural) encerra o parágrafo. O buffer acumula até o flush.
    buf: list[str] = []
    buf_kind: str | None = None

    def flush() -> None:
        nonlocal buf, buf_kind
        if not buf:
            return
        text = " ".join(buf)
        if buf_kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, text)
        elif buf_kind == "number":
            p = doc.add_paragraph(style="List Number")
            add_inline(p, text)
        elif buf_kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            add_inline(p, text)
            for run in p.runs:
                run.italic = True
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, text)
            stats["paragraphs"] += 1
        buf, buf_kind = [], None

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush()
            block, i = [], i + 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code_block(doc, block)
            stats["code"] += 1
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            if len(rows) >= 2:
                add_table(doc, rows)
                stats["tables"] += 1
            continue

        if stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            # O negrito já vem do estilo do título; os ** do Markdown sairiam literais.
            titulo = stripped[level:].strip().replace("**", "").replace("`", "")
            doc.add_heading(titulo, level=min(level, 4))
            stats["headings"] += 1

        elif stripped.startswith(("- ", "* ")):
            flush()
            buf, buf_kind = [stripped[2:]], "bullet"

        elif re.match(r"^\d+\.\s", stripped):
            flush()
            buf, buf_kind = [re.sub(r"^\d+\.\s", "", stripped)], "number"

        elif stripped.startswith(">"):
            content = stripped.lstrip("> ").strip()
            if buf_kind == "quote":
                buf.append(content)
            else:
                flush()
                buf, buf_kind = [content], "quote"

        elif stripped in ("---", "***", "___"):
            flush()

        elif not stripped:
            flush()

        elif buf_kind in ("bullet", "number") and line.startswith(("  ", "\t")):
            buf.append(stripped)  # continuação indentada do item de lista

        elif buf_kind in (None, "para"):
            buf.append(stripped)
            buf_kind = "para"

        else:
            flush()
            buf, buf_kind = [stripped], "para"

        i += 1

    flush()

    doc.save(docx_path)
    print(f"OK: {docx_path}")
    print(f"  títulos: {stats['headings']} | tabelas: {stats['tables']} | "
          f"blocos de código: {stats['code']} | parágrafos: {stats['paragraphs']}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
